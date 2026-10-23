import sqlite3
import json
import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

# Database path
db_path = os.path.join(os.path.dirname(__file__), 'data', 'schemes.db')
if not os.path.exists(db_path):
    db_path = os.path.join(os.path.dirname(__file__), '..', 'data', 'schemes.db')

conn = sqlite3.connect(db_path)
conn.row_factory = sqlite3.Row
cursor = conn.cursor()

cursor.execute("SELECT * FROM schemes ORDER BY scheme_id")
rows = cursor.fetchall()
print(f"Total schemes fetched: {len(rows)}")

wb = openpyxl.Workbook()
# Remove default sheet
wb.remove(wb.active)

# Header styling
header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
sub_header_fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid")

zebra_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

# -------------------------------------------------------------
# SHEET 1: All Schemes (Complete Human-Readable Dataset)
# -------------------------------------------------------------
ws1 = wb.create_sheet(title="TN Welfare Schemes")
ws1.views.sheetView[0].showGridLines = True

headers_sheet1 = [
    "S.No",
    "Scheme ID",
    "Scheme Name (English)",
    "Scheme Name (Tamil)",
    "Department",
    "Department (Tamil)",
    "Category",
    "Sub-Category",
    "Description (English)",
    "Description (Tamil)",
    "Official Benefits",
    "Eligibility Criteria",
    "Required Documents",
    "Application Steps",
    "Official Portal URLs",
    "Contact Helpline",
    "Status"
]

ws1.append(headers_sheet1)
for col_num in range(1, len(headers_sheet1) + 1):
    cell = ws1.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws1.row_dimensions[1].height = 28

def parse_json_safely(val, default):
    if not val:
        return default
    if isinstance(val, (dict, list)):
        return val
    try:
        return json.loads(val)
    except Exception:
        return default

def format_benefits_readable(b):
    b = parse_json_safely(b, {})
    if not b:
        return "Government Welfare Assistance"
    lines = []
    for k, v in b.items():
        title = k.replace("_", " ").title()
        lines.append(f"{title}: {v}")
    return "\n".join(lines)

def format_eligibility_readable(r):
    r = parse_json_safely(r, {})
    if not r:
        return "Open to all residents of Tamil Nadu"
    lines = []
    age = r.get("age")
    if age and isinstance(age, dict):
        min_a = age.get("min", 0)
        max_a = age.get("max", "No limit")
        lines.append(f"Age: {min_a} to {max_a} years")
    gender = r.get("gender")
    if gender and isinstance(gender, list) and "Any" not in gender:
        lines.append(f"Gender: {', '.join(gender)}")
    income = r.get("income_limit_annual")
    if income:
        lines.append(f"Annual Income Limit: Under Rs. {income:,}")
    schooling = r.get("schooling_requirement")
    if schooling:
        lines.append(f"Schooling: {schooling}")
    edu = r.get("education_level") or r.get("education")
    if edu:
        edu_str = ", ".join(edu) if isinstance(edu, list) else str(edu)
        lines.append(f"Education: {edu_str}")
    conditions = r.get("special_conditions", [])
    if conditions and isinstance(conditions, list):
        lines.append("Key Conditions:")
        for c in conditions:
            lines.append(f"  • {c}")
    return "\n".join(lines) if lines else "Tamil Nadu Resident"

def format_list_readable(val):
    items = parse_json_safely(val, [])
    if isinstance(items, list):
        return "\n".join([f"• {item}" for item in items])
    return str(val)

def format_steps_readable(val):
    items = parse_json_safely(val, [])
    if isinstance(items, list):
        return "\n".join(items)
    return str(val)

category_counts = {}

for idx, r in enumerate(rows, 1):
    row_dict = dict(r)
    cat = row_dict.get("category", "General Welfare")
    category_counts[cat] = category_counts.get(cat, 0) + 1

    benefits_txt = format_benefits_readable(row_dict.get("benefits"))
    elig_txt = format_eligibility_readable(row_dict.get("eligibility_rules"))
    docs_txt = format_list_readable(row_dict.get("required_documents"))
    steps_txt = format_steps_readable(row_dict.get("application_steps"))
    urls_txt = ", ".join(parse_json_safely(row_dict.get("official_urls"), []))

    row_data = [
        idx,
        row_dict.get("scheme_id", ""),
        row_dict.get("scheme_name_english", ""),
        row_dict.get("scheme_name_tamil", ""),
        row_dict.get("department", ""),
        row_dict.get("department_tamil", ""),
        cat,
        row_dict.get("sub_category", ""),
        row_dict.get("description_english", ""),
        row_dict.get("description_tamil", ""),
        benefits_txt,
        elig_txt,
        docs_txt,
        steps_txt,
        urls_txt,
        row_dict.get("contact_helpline", ""),
        row_dict.get("status", "Active")
    ]
    ws1.append(row_data)

    current_row = idx + 1
    row_fill = zebra_fill if idx % 2 == 0 else white_fill
    for col_num in range(1, len(row_data) + 1):
        c = ws1.cell(row=current_row, column=col_num)
        c.fill = row_fill
        c.border = thin_border
        c.font = Font(name="Calibri", size=10)
        # S.No and Scheme ID centered
        if col_num in [1, 2, 17]:
            c.alignment = Alignment(horizontal="center", vertical="top")
        else:
            c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

# -------------------------------------------------------------
# SHEET 2: Category Summary & Breakdown
# -------------------------------------------------------------
ws2 = wb.create_sheet(title="Category Statistics")
ws2.views.sheetView[0].showGridLines = True

cat_headers = ["S.No", "Welfare Category", "Number of Schemes", "Key Department Focus"]
ws2.append(cat_headers)
for col_num in range(1, len(cat_headers) + 1):
    cell = ws2.cell(row=1, column=col_num)
    cell.fill = sub_header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[1].height = 25

sorted_cats = sorted(category_counts.items(), key=lambda x: x[1], reverse=True)
for c_idx, (category_name, count) in enumerate(sorted_cats, 1):
    cat_row = [c_idx, category_name, count, "Tamil Nadu State Government"]
    ws2.append(cat_row)
    r_fill = zebra_fill if c_idx % 2 == 0 else white_fill
    for col_num in range(1, len(cat_row) + 1):
        c = ws2.cell(row=c_idx + 1, column=col_num)
        c.fill = r_fill
        c.border = thin_border
        c.font = Font(name="Calibri", size=11)
        if col_num in [1, 3]:
            c.alignment = Alignment(horizontal="center", vertical="center")
        else:
            c.alignment = Alignment(horizontal="left", vertical="center")

# Total summary row in Sheet 2
total_row = len(sorted_cats) + 2
ws2.cell(row=total_row, column=1, value="")
ws2.cell(row=total_row, column=2, value="TOTAL VERIFIED SCHEMES")
ws2.cell(row=total_row, column=3, value=len(rows))
ws2.cell(row=total_row, column=4, value="All Departments")
for col_num in range(1, 5):
    c = ws2.cell(row=total_row, column=col_num)
    c.font = Font(name="Calibri", size=11, bold=True, color="065F46")
    c.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    c.border = thin_border

# -------------------------------------------------------------
# SHEET 3: Raw JSON Columns (For Developers & Data Engineers)
# -------------------------------------------------------------
ws3 = wb.create_sheet(title="Raw JSON Data")
ws3.views.sheetView[0].showGridLines = True

headers_sheet3 = [
    "Scheme ID",
    "Scheme Name (English)",
    "Benefits (JSON)",
    "Eligibility Rules (JSON)",
    "Required Documents (JSON)",
    "Application Steps (JSON)",
    "Official URLs (JSON)",
    "Contact Helpline"
]
ws3.append(headers_sheet3)
for col_num in range(1, len(headers_sheet3) + 1):
    cell = ws3.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[1].height = 25

for idx, r in enumerate(rows, 1):
    row_dict = dict(r)
    r_data = [
        row_dict.get("scheme_id", ""),
        row_dict.get("scheme_name_english", ""),
        row_dict.get("benefits", ""),
        row_dict.get("eligibility_rules", ""),
        row_dict.get("required_documents", ""),
        row_dict.get("application_steps", ""),
        row_dict.get("official_urls", ""),
        row_dict.get("contact_helpline", "")
    ]
    ws3.append(r_data)
    r_fill = zebra_fill if idx % 2 == 0 else white_fill
    for col_num in range(1, len(r_data) + 1):
        c = ws3.cell(row=idx + 1, column=col_num)
        c.fill = r_fill
        c.border = thin_border
        c.font = Font(name="Consolas", size=9)
        if col_num == 1:
            c.alignment = Alignment(horizontal="center", vertical="top")
        else:
            c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

# -------------------------------------------------------------
# Adjust Column Widths across all sheets
# -------------------------------------------------------------
col_widths_sheet1 = {
    1: 8,    # S.No
    2: 24,   # Scheme ID
    3: 38,   # Name EN
    4: 38,   # Name TA
    5: 30,   # Dept
    6: 30,   # Dept TA
    7: 28,   # Category
    8: 24,   # Sub-category
    9: 45,   # Desc EN
    10: 45,  # Desc TA
    11: 45,  # Benefits
    12: 45,  # Eligibility
    13: 40,  # Documents
    14: 45,  # Steps
    15: 35,  # URLs
    16: 28,  # Helpline
    17: 12   # Status
}

for col_num, width in col_widths_sheet1.items():
    ws1.column_dimensions[get_column_letter(col_num)].width = width

ws2.column_dimensions['A'].width = 10
ws2.column_dimensions['B'].width = 45
ws2.column_dimensions['C'].width = 22
ws2.column_dimensions['D'].width = 35

ws3.column_dimensions['A'].width = 24
ws3.column_dimensions['B'].width = 35
ws3.column_dimensions['C'].width = 50
ws3.column_dimensions['D'].width = 50
ws3.column_dimensions['E'].width = 45
ws3.column_dimensions['F'].width = 50
ws3.column_dimensions['G'].width = 35
ws3.column_dimensions['H'].width = 25

# Save files to multiple convenient locations
target_files = [
    os.path.join(os.path.dirname(__file__), "TamilNadu_Government_Schemes_Dataset.xlsx"),
    os.path.join(os.path.dirname(__file__), "..", "TamilNadu_Government_Schemes_Dataset.xlsx"),
    os.path.join(os.path.dirname(__file__), "data", "TamilNadu_Government_Schemes_Dataset.xlsx")
]

for out_file in target_files:
    try:
        os.makedirs(os.path.dirname(out_file), exist_ok=True)
        wb.save(out_file)
        print(f"Saved Excel file successfully: {os.path.abspath(out_file)}")
    except Exception as e:
        print(f"Could not save to {out_file}: {e}")

# Also generate a clean CSV file
import csv
csv_out_path = os.path.join(os.path.dirname(__file__), "..", "TamilNadu_Government_Schemes_Dataset.csv")
with open(csv_out_path, mode="w", encoding="utf-8-sig", newline="") as f:
    writer = csv.writer(f)
    writer.writerow(headers_sheet1)
    for idx, r in enumerate(rows, 1):
        row_dict = dict(r)
        cat = row_dict.get("category", "")
        writer.writerow([
            idx,
            row_dict.get("scheme_id", ""),
            row_dict.get("scheme_name_english", ""),
            row_dict.get("scheme_name_tamil", ""),
            row_dict.get("department", ""),
            row_dict.get("department_tamil", ""),
            cat,
            row_dict.get("sub_category", ""),
            row_dict.get("description_english", ""),
            row_dict.get("description_tamil", ""),
            format_benefits_readable(row_dict.get("benefits")),
            format_eligibility_readable(row_dict.get("eligibility_rules")),
            format_list_readable(row_dict.get("required_documents")),
            format_steps_readable(row_dict.get("application_steps")),
            ", ".join(parse_json_safely(row_dict.get("official_urls"), [])),
            row_dict.get("contact_helpline", ""),
            row_dict.get("status", "Active")
        ])
print(f"Saved CSV file successfully: {os.path.abspath(csv_out_path)}")
